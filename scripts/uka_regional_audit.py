"""Kontrollera UKÄ:s regionala tabellbilagor utan att publicera osäkra data."""
import csv
import json
import re
from pathlib import Path
from urllib.parse import urljoin
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

PAGE = "https://www.uka.se/statistik-och-analys/om-var-statistik/statistiska-meddelanden"
OUT = Path("output/uka_regional_audit")
OUT.mkdir(parents=True, exist_ok=True)
HEADERS = {"User-Agent": "Kompetensarena-Norrbotten-UKA-audit/1.0"}
KEYWORDS = ("rekryteringslän", "lärosätets län", "bosatta", "bosatt", "regional rekrytering", "examinerade")
MAX_MATCHES = 300
session = requests.Session()
session.headers.update(HEADERS)

def fetch(url):
    response = session.get(url, timeout=180)
    response.raise_for_status()
    return response

def main():
    page = BeautifulSoup(fetch(PAGE).text, "html.parser")
    links = []
    for anchor in page.select("a[href]"):
        name = anchor.get_text(" ", strip=True)
        if "Tabellbilaga till Studenter och examinerade" in name and "2024/25" in name:
            links.append((name, urljoin(PAGE, anchor["href"])))
    if len(links) != 1:
        raise RuntimeError(f"Förväntade en UF0205-bilaga men hittade {len(links)}: {links}")
    name, url = links[0]
    source = OUT / "uf0205_current.xlsx"
    source.write_bytes(fetch(url).content)
    workbook = load_workbook(source, read_only=True, data_only=True)
    report = {"source_page": PAGE, "source_url": url, "source_title": name, "sheets": [], "matches_truncated_at": MAX_MATCHES}
    matches = []
    for sheet in workbook:
        sheet.reset_dimensions()
        sheet_report = {"name": sheet.title, "hits": 0}
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(v).strip() if v is not None else "" for v in row]
            found = [term for term in KEYWORDS if term in " ".join(values).lower()]
            if found:
                sheet_report["hits"] += 1
                if len(matches) < MAX_MATCHES:
                    matches.append({"sheet": sheet.title, "row": row_no, "matched": "|".join(found), "cells": " | ".join(v for v in values if v)[:2500]})
        report["sheets"].append(sheet_report)
    workbook.close()
    (OUT / "inventory.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT / "candidate_rows.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["sheet", "row", "matched", "cells"])
        writer.writeheader()
        writer.writerows(matches)
    print(json.dumps({"source": url, "sheets": report["sheets"], "candidate_rows": len(matches)}, ensure_ascii=False, indent=2))
    if not matches:
        raise RuntimeError("Hittade inga regionala kandidatfält: kontrollera källstruktur innan import.")

if __name__ == "__main__":
    main()
